from __future__ import annotations

import io
import json
import sys
import unittest
import zipfile
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import docx

from app.resume_intelligence.parser import (
    ResumeParseError,
    analyze_extracted_document,
    build_resume_summary,
    extract_document,
    parse_resume_analysis,
    to_legacy_parsed_data,
)


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for value in paragraphs:
        document.add_paragraph(value)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(lines: list[str]) -> bytes:
    def escape(value: str) -> str:
        return value.replace(chr(92), chr(92) * 2).replace("(", r"\(").replace(")", r"\)")

    text_operations = ["BT", "/F1 12 Tf", "72 720 Td"]
    for index, line in enumerate(lines):
        if index:
            text_operations.append("0 -18 Td")
        text_operations.append(f"({escape(line)}) Tj")
    text_operations.append("ET")
    stream = "\n".join(text_operations).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(output)


RESUME_PARAGRAPHS = [
    "John Doe",
    "john@example.com | +91 98765 43210",
    "TECHNICAL SKILLS",
    "Python, React, ReactJS, SQL, C++",
    "PROJECTS",
    "Campus Portal",
    "Built a React dashboard with a Python API, reducing load time by 30%.",
    "EXPERIENCE",
    "Software Intern at Acme Ltd",
    "Jan 2024 - Jun 2024",
    "Developed FastAPI services and PostgreSQL queries.",
    "EDUCATION",
    "B.Tech in Computer Science",
    "ABC University",
    "2021 - 2025",
    "CERTIFICATIONS",
    "AWS Certified Cloud Practitioner - Amazon",
]


class DocumentExtractionTests(unittest.TestCase):
    def test_docx_text_and_tables_are_extracted(self) -> None:
        content = make_docx(["TECHNICAL SKILLS"], [["Python", "React"]])
        document = extract_document("resume.docx", None, content)
        self.assertIn("TECHNICAL SKILLS", document.text)
        self.assertIn("Python | React", document.text)
        self.assertEqual(len(document.sha256), 64)
        self.assertIsNone(document.page_count)

    def test_pdf_text_is_extracted(self) -> None:
        content = make_pdf(["TECHNICAL SKILLS", "Python and React"])
        document = extract_document("resume.pdf", "application/pdf", content)
        self.assertIn("Python and React", document.text)
        self.assertEqual(document.page_count, 1)

    def test_scanned_pdf_reports_ocr_requirement(self) -> None:
        with self.assertRaises(ResumeParseError) as raised:
            extract_document("scan.pdf", "application/pdf", make_pdf([]))
        self.assertEqual(raised.exception.status_code, 422)
        self.assertEqual(raised.exception.code, "ocr_required")
        self.assertIn("OCR", raised.exception.message)

    def test_corrupt_pdf_is_rejected(self) -> None:
        with self.assertRaises(ResumeParseError) as raised:
            extract_document(
                "broken.pdf", "application/pdf", b"%PDF-1.4\nnot a valid pdf"
            )
        self.assertEqual(raised.exception.status_code, 422)

    def test_empty_and_corrupt_docx_are_rejected(self) -> None:
        with self.assertRaises(ResumeParseError) as empty:
            extract_document("empty.docx", None, make_docx([]))
        self.assertEqual(empty.exception.code, "empty_document")
        with self.assertRaises(ResumeParseError) as corrupt:
            extract_document("broken.docx", None, b"PK\x03\x04broken")
        self.assertEqual(corrupt.exception.code, "invalid_docx")

    def test_macro_enabled_document_is_rejected(self) -> None:
        source = make_docx(["TECHNICAL SKILLS", "Python"])
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as source_archive:
            with zipfile.ZipFile(output, "w") as target_archive:
                for entry in source_archive.infolist():
                    target_archive.writestr(entry, source_archive.read(entry.filename))
                target_archive.writestr("word/vbaProject.bin", b"macro")
        with self.assertRaises(ResumeParseError) as raised:
            extract_document("resume.docx", None, output.getvalue())
        self.assertEqual(raised.exception.code, "macro_enabled_document")

    def test_file_type_extension_signature_and_size_are_validated(self) -> None:
        with self.assertRaises(ResumeParseError) as unsupported:
            extract_document("resume.txt", "text/plain", b"hello")
        self.assertEqual(unsupported.exception.status_code, 415)
        with self.assertRaises(ResumeParseError) as mismatch:
            extract_document("resume.pdf", "image/png", make_pdf(["Python"]))
        self.assertEqual(mismatch.exception.code, "content_type_mismatch")
        with self.assertRaises(ResumeParseError) as signature:
            extract_document("resume.pdf", "application/pdf", b"not a pdf")
        self.assertEqual(signature.exception.code, "invalid_file_signature")
        with self.assertRaises(ResumeParseError) as too_large:
            extract_document("resume.pdf", "application/pdf", b"%PDF-" + b"x" * 30, max_bytes=10)
        self.assertEqual(too_large.exception.status_code, 413)
        with self.assertRaises(ResumeParseError) as empty:
            extract_document("resume.pdf", "application/pdf", b"")
        self.assertEqual(empty.exception.code, "empty_file")


class StructuredAnalysisTests(unittest.TestCase):
    def test_sections_skills_and_evidence_are_structured(self) -> None:
        analysis = parse_resume_analysis(
            "resume.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            make_docx(RESUME_PARAGRAPHS),
        )
        section_types = [section.section_type for section in analysis.sections]
        self.assertTrue(
            {"header", "skills", "projects", "experience", "education", "certifications"}
            <= set(section_types)
        )
        skills = {skill.skill_id: skill for skill in analysis.skills}
        self.assertEqual(skills["react"].mention_count, 3)
        self.assertEqual(skills["react"].evidence_strength, "demonstrated_outcome")
        self.assertEqual(skills["sql"].evidence_strength, "listed")
        self.assertEqual(skills["sql"].evidence_type, "explicit_mention")
        self.assertEqual(skills["python"].section_type, "projects")
        self.assertEqual(analysis.stats.project_count, 1)
        self.assertEqual(analysis.stats.experience_count, 1)
        self.assertEqual(analysis.stats.education_count, 1)
        self.assertEqual(analysis.stats.certification_count, 1)

    def test_projects_experience_education_and_certifications_are_not_fabricated(self) -> None:
        analysis = parse_resume_analysis("resume.docx", None, make_docx(RESUME_PARAGRAPHS))
        project = analysis.projects[0]
        experience = analysis.work_experience[0]
        education = analysis.education[0]
        certification = analysis.certifications[0]
        self.assertEqual(project.title, "Campus Portal")
        self.assertIn("reducing load time by 30%", project.description)
        self.assertEqual(experience.role, "Software Intern")
        self.assertEqual(experience.organization, "Acme Ltd")
        self.assertEqual(experience.duration, "Jan 2024 - Jun 2024")
        self.assertEqual(education.institution, "ABC University")
        self.assertEqual(education.degree, "B.Tech in Computer Science")
        self.assertEqual(education.year, "2021 - 2025")
        self.assertEqual(certification.name, "AWS Certified Cloud Practitioner")
        self.assertEqual(certification.issuer, "Amazon")

    def test_contact_values_are_not_retained_in_analysis(self) -> None:
        analysis = parse_resume_analysis("resume.docx", None, make_docx(RESUME_PARAGRAPHS))
        serialized = json.dumps(analysis.model_dump(mode="json"))
        self.assertNotIn("john@example.com", serialized)
        self.assertNotIn("98765", serialized)
        self.assertTrue(analysis.contact.email_present)
        self.assertTrue(analysis.contact.phone_present)
        self.assertNotIn("john@example.com", build_resume_summary(analysis))

    def test_contact_heading_is_accepted_and_masked(self) -> None:
        analysis = parse_resume_analysis(
            "resume.docx",
            None,
            make_docx(
                [
                    "CONTACT",
                    "jane@example.com | +91 98765 43210",
                    "TECHNICAL SKILLS",
                    "Python",
                ]
            ),
        )
        contact_sections = [
            section for section in analysis.sections if section.section_type == "contact"
        ]
        self.assertEqual(len(contact_sections), 1)
        self.assertTrue(analysis.contact.contact_section_present)
        self.assertTrue(analysis.contact.email_present)
        self.assertNotIn("jane@example.com", json.dumps(analysis.model_dump(mode="json")))

    def test_aliases_are_grouped_without_double_counting_skills(self) -> None:
        analysis = parse_resume_analysis(
            "resume.docx",
            None,
            make_docx(["TECHNICAL SKILLS", "React, ReactJS, React.js"]),
        )
        self.assertEqual(len(analysis.skills), 1)
        self.assertEqual(analysis.skills[0].skill_id, "react")
        self.assertEqual(analysis.skills[0].mention_count, 3)
        self.assertEqual(len({skill.skill_id for skill in analysis.skills}), len(analysis.skills))

    def test_source_spans_point_to_the_matched_alias(self) -> None:
        content = make_docx(["PROJECTS", "Built a Python API"])
        document = extract_document("resume.docx", None, content)
        analysis = analyze_extracted_document(document)
        python_skill = next(skill for skill in analysis.skills if skill.skill_id == "python")
        span_text = document.text[python_skill.span.start:python_skill.span.end]
        self.assertEqual(span_text.casefold(), python_skill.matched_alias.casefold())

    def test_empty_structured_sections_do_not_create_placeholder_records(self) -> None:
        analysis = parse_resume_analysis(
            "resume.docx", None, make_docx(["SUMMARY", "Computer science student."])
        )
        self.assertEqual(analysis.projects, [])
        self.assertEqual(analysis.work_experience, [])
        self.assertEqual(analysis.education, [])
        self.assertEqual(analysis.certifications, [])
        self.assertEqual(analysis.skills, [])
        self.assertIn("No catalog skills", " ".join(analysis.warnings))

    def test_legacy_shape_preserves_compatibility_without_placeholders(self) -> None:
        analysis = parse_resume_analysis("resume.docx", None, make_docx(RESUME_PARAGRAPHS))
        parsed = to_legacy_parsed_data(analysis)
        self.assertIn("React", parsed["skills"])
        self.assertEqual(parsed["projects"][0]["name"], "Campus Portal")
        self.assertEqual(parsed["education"][0]["institution"], "ABC University")
        self.assertEqual(parsed["work_experience"][0]["company"], "Acme Ltd")


if __name__ == "__main__":
    unittest.main()
