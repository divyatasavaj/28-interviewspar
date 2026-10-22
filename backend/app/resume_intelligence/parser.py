"""Safe resume extraction and deterministic, evidence-backed skill parsing."""

from __future__ import annotations

import bisect
import hashlib
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from .catalog import ResumeIntelligenceCatalog, load_catalog
from .normalization import SkillMatch, SkillNormalizer
from .schemas import (
    CertificationEvidence,
    ContactMetadata,
    EducationEvidence,
    ExtractionStats,
    FileMetadata,
    ProjectEvidence,
    ResumeAnalysis,
    ResumeSection,
    SkillEvidence,
    SourceSpan,
    WorkExperienceEvidence,
)


PARSER_VERSION = "resume-parser-1.0.0"
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 1_000_000
MAX_PDF_PAGES = 50
MAX_DOCX_ENTRIES = 10_000
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

ALLOWED_EXTENSIONS = frozenset({".pdf", ".docx"})
EXPECTED_MEDIA_TYPES = {
    ".pdf": frozenset({"application/pdf"}),
    ".docx": frozenset(
        {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        }
    ),
}
GENERIC_MEDIA_TYPES = frozenset({"application/octet-stream", "binary/octet-stream"})

_EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
_PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d[\d ()-]{7,}\d)(?!\d)")
_PROFILE_LINK_RE = re.compile(
    r"(?:https?://|www\.|linkedin\.com/|github\.com/)", re.IGNORECASE
)
_BULLET_RE = re.compile(r"^\s*(?:[-*•●▪◦]|\d+[.)])\s+")
_OUTCOME_RE = re.compile(
    r"\b\d+(?:\.\d+)?\s*(?:%|percent|ms|seconds?|minutes?|hours?|days?|weeks?|"
    r"months?|users?|records?|requests?|transactions?|x)(?!\w)",
    re.IGNORECASE,
)
_ACTION_VERBS = frozenset(
    {
        "analyzed",
        "applied",
        "architected",
        "automated",
        "built",
        "created",
        "deployed",
        "designed",
        "developed",
        "engineered",
        "implemented",
        "improved",
        "integrated",
        "launched",
        "led",
        "maintained",
        "migrated",
        "optimized",
        "reduced",
        "scaled",
        "tested",
    }
)
_DEGREE_RE = re.compile(
    r"\b(?:B\.?\s?Tech|B\.?\s?E\.?|B\.?\s?Sc\.?|B\.?\s?S\.?|"
    r"Bachelor(?:'s)?(?:\s+of\s+[A-Za-z ]{2,40})?|M\.?\s?Tech|M\.?\s?E\.?|"
    r"M\.?\s?Sc\.?|M\.?\s?S\.?|Master(?:'s)?(?:\s+of\s+[A-Za-z ]{2,40})?|"
    r"MCA|BCA|Diploma|Ph\.?\s?D\.?)\b[^,|\n;]{0,100}",
    re.IGNORECASE,
)
_INSTITUTION_RE = re.compile(
    r"(?:[A-Z][A-Za-z&.'-]*\s+){0,1}(?:University|College|Institute|IIT|NIT|"
    r"Academy)(?:\s+of\s+(?:[A-Z][A-Za-z&.'-]*\s*){1,3})?",
    re.IGNORECASE,
)
_YEAR_RE = re.compile(
    r"\b(?:(?:19|20)\d{2}\s*(?:-|–|—|to)\s*(?:(?:19|20)\d{2}|present|current)|"
    r"(?:19|20)\d{2})\b",
    re.IGNORECASE,
)
_DATE_RANGE_RE = re.compile(
    r"\b(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+"
    r"(?:19|20)\d{2}|(?:19|20)\d{2})\s*(?:-|–|—|to)\s*"
    r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+"
    r"(?:19|20)\d{2}|present|current|(?:19|20)\d{2})\b",
    re.IGNORECASE,
)
_COMPANY_HINT_RE = re.compile(
    r"\b(?:Ltd|Limited|Inc|Corp|Corporation|Pvt|Private|Technologies|Technology|"
    r"Solutions|Systems|Labs|Company|Enterprises|Services)\b",
    re.IGNORECASE,
)


class ResumeParseError(ValueError):
    """A user-safe resume parsing failure with an API status code."""

    def __init__(self, message: str, status_code: int, code: str):
        super().__init__(message)
        self.message = message
        self.status_code = status_code
        self.code = code


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    extension: str
    media_type: str
    size_bytes: int
    sha256: str
    page_count: int | None


@dataclass(frozen=True)
class _Line:
    text: str
    start: int
    end: int
    index: int


@dataclass(frozen=True)
class _Section:
    section_id: str
    section_type: str
    heading: str | None
    heading_line_index: int | None
    start: int
    end: int
    content_start: int
    line_start: int
    line_end: int


@dataclass(frozen=True)
class _Block:
    lines: tuple[_Line, ...]
    start: int
    end: int


@dataclass(frozen=True)
class _Occurrence:
    match: SkillMatch
    section: _Section
    snippet: str
    evidence_type: str
    evidence_strength: str
    rank: int


@lru_cache(maxsize=1)
def _catalog() -> ResumeIntelligenceCatalog:
    return load_catalog()


@lru_cache(maxsize=1)
def _normalizer() -> SkillNormalizer:
    return SkillNormalizer(_catalog().skills.values())


def extract_document(
    filename: str | None,
    content_type: str | None,
    content: bytes,
    *,
    max_bytes: int = MAX_FILE_BYTES,
) -> ExtractedDocument:
    """Validate and extract text without executing or persisting the upload."""

    if not content:
        raise ResumeParseError("The uploaded file is empty.", 422, "empty_file")
    if len(content) > max_bytes:
        raise ResumeParseError(
            f"File too large (max {max_bytes // (1024 * 1024)}MB).",
            413,
            "file_too_large",
        )

    safe_name = (filename or "").replace(chr(92), "/").rsplit("/", 1)[-1]
    extension = Path(safe_name).suffix.lower()
    if not safe_name or extension not in ALLOWED_EXTENSIONS:
        raise ResumeParseError(
            "Only PDF and DOCX files are supported.", 415, "unsupported_file_type"
        )

    declared_type = (content_type or "").split(";", 1)[0].strip().lower()
    if (
        declared_type
        and declared_type not in EXPECTED_MEDIA_TYPES[extension]
        and declared_type not in GENERIC_MEDIA_TYPES
    ):
        raise ResumeParseError(
            "The declared file type does not match the file extension.",
            415,
            "content_type_mismatch",
        )

    if extension == ".pdf":
        if b"%PDF-" not in content[:1024]:
            raise ResumeParseError(
                "The file does not contain valid PDF content.",
                415,
                "invalid_file_signature",
            )
        text, page_count = _extract_pdf(content)
        media_type = "application/pdf"
    else:
        _validate_docx_archive(content)
        text = _extract_docx(content)
        page_count = None
        media_type = next(iter(EXPECTED_MEDIA_TYPES[extension]))

    cleaned = _clean_text(text)
    if not cleaned:
        if extension == ".pdf":
            raise ResumeParseError(
                "This PDF contains no extractable text. It may be a scanned image; "
                "OCR is required before resume analysis.",
                422,
                "ocr_required",
            )
        raise ResumeParseError(
            "The DOCX file contains no extractable text.", 422, "empty_document"
        )
    if len(cleaned) > MAX_EXTRACTED_CHARACTERS:
        raise ResumeParseError(
            "The extracted resume text is too large to analyze.",
            413,
            "extracted_text_too_large",
        )

    return ExtractedDocument(
        text=cleaned,
        extension=extension,
        media_type=media_type,
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        page_count=page_count,
    )


def _extract_pdf(content: bytes) -> tuple[str, int]:
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            page_count = len(pdf.pages)
            if page_count > MAX_PDF_PAGES:
                raise ResumeParseError(
                    f"PDF has too many pages (max {MAX_PDF_PAGES}).",
                    422,
                    "too_many_pages",
                )
            page_text = []
            for page in pdf.pages:
                extracted = page.extract_text() or ""
                if extracted.strip():
                    page_text.append(extracted)
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(
            "The PDF could not be read. It may be encrypted, corrupted, or malformed.",
            422,
            "pdf_read_error",
        ) from exc

    text = "\n\n".join(page_text).strip()
    if sum(character.isalnum() for character in text) < 20:
        raise ResumeParseError(
            "This PDF contains no extractable text. It may be a scanned image; "
            "OCR is required before resume analysis.",
            422,
            "ocr_required",
        )
    return text, page_count


def _validate_docx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise ResumeParseError(
                    "The DOCX archive contains too many entries.", 422, "docx_too_complex"
                )
            total_uncompressed = sum(entry.file_size for entry in entries)
            if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ResumeParseError(
                    "The DOCX archive is too large after decompression.",
                    413,
                    "docx_decompression_limit",
                )
            names = {entry.filename.lower() for entry in entries}
            if "word/document.xml" not in names or "[content_types].xml" not in names:
                raise ResumeParseError(
                    "The file does not contain a valid DOCX document.",
                    415,
                    "invalid_docx",
                )
            if any(name.endswith("vbaproject.bin") for name in names):
                raise ResumeParseError(
                    "Macro-enabled Word documents are not supported. Save as DOCX first.",
                    415,
                    "macro_enabled_document",
                )
    except ResumeParseError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise ResumeParseError(
            "The file does not contain valid DOCX content.",
            415,
            "invalid_docx",
        ) from exc


def _extract_docx(content: bytes) -> str:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(io.BytesIO(content))
        parts: list[str] = []
        if hasattr(document, "iter_inner_content"):
            for item in document.iter_inner_content():
                if isinstance(item, Paragraph):
                    parts.append(item.text)
                elif isinstance(item, Table):
                    for row in item.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        else:
            parts.extend(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise ResumeParseError(
            "The DOCX file could not be read. It may be corrupted or malformed.",
            422,
            "docx_read_error",
        ) from exc


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x0b", "\n").replace("\x0c", "\n").replace("\x85", "\n")
    text = "".join(
        character
        for character in text
        if character in "\n\t" or not unicodedata.category(character).startswith("C")
    )
    output: list[str] = []
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped:
            output.append(stripped)
        elif output and output[-1] != "":
            output.append("")
    while output and output[-1] == "":
        output.pop()
    return "\n".join(output).strip()


_HEADING_MAP = {
    "contact": "contact",
    "contact information": "contact",
    "contact details": "contact",
    "personal details": "contact",
    "summary": "summary",
    "professional summary": "summary",
    "profile": "summary",
    "professional profile": "summary",
    "objective": "summary",
    "career objective": "summary",
    "about": "summary",
    "about me": "summary",
    "skills": "skills",
    "technical skills": "skills",
    "core skills": "skills",
    "key skills": "skills",
    "skill set": "skills",
    "technical expertise": "skills",
    "technical competencies": "skills",
    "core competencies": "skills",
    "technologies": "skills",
    "tools and technologies": "skills",
    "projects": "projects",
    "project": "projects",
    "academic projects": "projects",
    "personal projects": "projects",
    "key projects": "projects",
    "project experience": "projects",
    "experience": "experience",
    "work experience": "experience",
    "professional experience": "experience",
    "employment": "experience",
    "employment history": "experience",
    "internship": "experience",
    "internships": "experience",
    "education": "education",
    "academic background": "education",
    "academics": "education",
    "academic qualifications": "education",
    "qualifications": "education",
    "certifications": "certifications",
    "certification": "certifications",
    "certificates": "certifications",
    "licenses": "certifications",
    "certifications and courses": "certifications",
    "awards": "awards",
    "achievements": "awards",
    "honors": "awards",
    "publications": "publications",
    "languages": "languages",
    "interests": "interests",
    "hobbies": "interests",
}


def _heading_key(value: str) -> str:
    stripped = value.strip().strip(":").strip()
    key = re.sub(r"[^a-z0-9&]+", " ", stripped.casefold())
    return re.sub(r"\s+", " ", key).strip()


def _classify_heading(value: str) -> str | None:
    stripped = value.strip()
    if not stripped or len(stripped) > 60 or len(stripped.split()) > 8:
        return None
    if any(mark in stripped for mark in (".", "!", "?", ";", ",")):
        return None
    key = _heading_key(stripped)
    if key in _HEADING_MAP:
        return _HEADING_MAP[key]
    if re.fullmatch(r"(?:technical|core|key)\s+skills(?:\s*(?:&|and)\s*tools)?", key):
        return "skills"
    if re.fullmatch(r"(?:academic|personal|key)?\s*projects", key):
        return "projects"
    if re.fullmatch(
        r"(?:work|professional|employment|industry)?\s*experience(?:\s+history)?",
        key,
    ):
        return "experience"
    if re.fullmatch(r"(?:academic|educational)\s+(?:background|qualifications)", key):
        return "education"
    return None


def _make_lines(text: str) -> tuple[_Line, ...]:
    lines: list[_Line] = []
    offset = 0
    for index, value in enumerate(text.split("\n")):
        lines.append(_Line(text=value, start=offset, end=offset + len(value), index=index))
        offset += len(value) + 1
    return tuple(lines)


def _parse_sections(text: str, lines: tuple[_Line, ...]) -> tuple[_Section, ...]:
    headings: list[tuple[int, str, _Line]] = []
    for line in lines:
        section_type = _classify_heading(line.text)
        if section_type is not None:
            headings.append((line.index, section_type, line))

    sections: list[_Section] = []
    if not headings:
        if text.strip():
            sections.append(
                _Section(
                    section_id="section-1",
                    section_type="other",
                    heading=None,
                    heading_line_index=None,
                    start=0,
                    end=len(text),
                    content_start=0,
                    line_start=0,
                    line_end=max(0, len(lines) - 1),
                )
            )
        return tuple(sections)

    first_heading_index = headings[0][0]
    if any(line.text.strip() for line in lines[:first_heading_index]):
        sections.append(
            _Section(
                section_id="section-1",
                section_type="header",
                heading=None,
                heading_line_index=None,
                start=0,
                end=headings[0][2].start,
                content_start=0,
                line_start=0,
                line_end=max(0, first_heading_index - 1),
            )
        )

    for position, (line_index, section_type, heading_line) in enumerate(headings):
        next_start = (
            headings[position + 1][2].start if position + 1 < len(headings) else len(text)
        )
        next_line_index = (
            headings[position + 1][0]
            if position + 1 < len(headings)
            else max(0, len(lines) - 1)
        )
        sections.append(
            _Section(
                section_id=f"section-{len(sections) + 1}",
                section_type=section_type,
                heading=heading_line.text.strip().strip(":").strip(),
                heading_line_index=line_index,
                start=heading_line.start,
                end=next_start,
                content_start=heading_line.end,
                line_start=line_index,
                line_end=next_line_index if next_start == len(text) else next_line_index - 1,
            )
        )
    return tuple(sections)


def _section_for_offset(
    sections: tuple[_Section, ...], offset: int
) -> _Section | None:
    starts = [section.start for section in sections]
    position = bisect.bisect_right(starts, offset) - 1
    if position < 0:
        return None
    section = sections[position]
    if offset < section.content_start or offset >= section.end:
        return None
    return section


def _line_for_offset(lines: tuple[_Line, ...], offset: int) -> _Line:
    starts = [line.start for line in lines]
    position = bisect.bisect_right(starts, offset) - 1
    return lines[max(0, position)]


def _extract_skill_evidence(
    text: str,
    lines: tuple[_Line, ...],
    sections: tuple[_Section, ...],
    catalog: ResumeIntelligenceCatalog,
    normalizer: SkillNormalizer,
) -> list[SkillEvidence]:
    occurrences: dict[str, list[_Occurrence]] = {}
    for match in normalizer.normalize(text):
        section = _section_for_offset(sections, match.start)
        if section is None or section.section_type == "contact":
            continue
        line = _line_for_offset(lines, match.start)
        evidence_type, evidence_strength, rank = _classify_evidence(
            section.section_type, line.text
        )
        snippet = _evidence_snippet(text, line, match)
        occurrence = _Occurrence(
            match=match,
            section=section,
            snippet=snippet,
            evidence_type=evidence_type,
            evidence_strength=evidence_strength,
            rank=rank,
        )
        occurrences.setdefault(match.skill_id, []).append(occurrence)

    evidence: list[SkillEvidence] = []
    for skill_id, items in occurrences.items():
        best = max(
            items,
            key=lambda item: (
                item.rank,
                item.match.confidence,
                len(item.match.matched_alias),
            ),
        )
        section_types_seen: list[str] = []
        for item in items:
            if item.section.section_type not in section_types_seen:
                section_types_seen.append(item.section.section_type)
        evidence.append(
            SkillEvidence(
                skill_id=skill_id,
                canonical_name=best.match.canonical_name,
                matched_alias=best.match.matched_alias,
                mention_count=len(items),
                section_id=best.section.section_id,
                section_type=best.section.section_type,
                section_types_seen=section_types_seen,
                evidence_type=best.evidence_type,
                evidence_strength=best.evidence_strength,
                evidence_snippet=best.snippet,
                span=SourceSpan(start=best.match.start, end=best.match.end),
                extraction_method=best.match.method,
                catalog_version=catalog.skills_catalog_version,
                confidence=round(best.match.confidence, 3),
                context_validated=best.match.context_validated,
            )
        )
    return sorted(evidence, key=lambda item: item.canonical_name.casefold())


def _classify_evidence(section_type: str, line_text: str) -> tuple[str, str, int]:
    if section_type == "skills":
        return "explicit_mention", "listed", 1
    if section_type in {"projects", "experience"}:
        words = {word.casefold() for word in re.findall(r"[A-Za-z]+", line_text)}
        has_action = bool(words & _ACTION_VERBS)
        has_outcome = bool(_OUTCOME_RE.search(line_text))
        if has_action and has_outcome:
            return "contextual_evidence", "demonstrated_outcome", 4
        if has_action:
            return "contextual_evidence", "demonstrated_use", 3
        return "contextual_evidence", "contextual_mention", 2
    return "contextual_evidence", "contextual_mention", 2


def _evidence_snippet(text: str, line: _Line, match: SkillMatch) -> str:
    snippet = _BULLET_RE.sub("", line.text.strip()).strip()
    snippet = _EMAIL_RE.sub("[redacted-email]", snippet)
    snippet = _PROFILE_LINK_RE.sub("[profile-link]", snippet)
    snippet = _PHONE_RE.sub(_redact_phone_match, snippet)
    if len(snippet) <= 240:
        return snippet or text[match.start:match.end]

    relative_start = max(0, match.start - line.start)
    crop_start = max(0, relative_start - 80)
    crop_end = min(len(snippet), crop_start + 220)
    if crop_end == len(snippet):
        crop_start = max(0, crop_end - 220)
    prefix = "..." if crop_start > 0 else ""
    suffix = "..." if crop_end < len(snippet) else ""
    return f"{prefix}{snippet[crop_start:crop_end].strip()}{suffix}"[:240]


def _contains_phone(text: str) -> bool:
    return any(
        sum(character.isdigit() for character in match.group()) >= 10
        for match in _PHONE_RE.finditer(text)
    )


def _redact_phone_match(match: re.Match[str]) -> str:
    return (
        "[redacted-phone]"
        if sum(character.isdigit() for character in match.group()) >= 10
        else match.group()
    )


def _redact_sensitive(text: str) -> str:
    text = _EMAIL_RE.sub("[redacted-email]", text)
    text = _PROFILE_LINK_RE.sub("[profile-link]", text)
    return _PHONE_RE.sub(_redact_phone_match, text)


def _section_content_lines(
    section: _Section, lines: tuple[_Line, ...]
) -> tuple[_Line, ...]:
    return tuple(
        line
        for line in lines
        if line.start >= section.content_start
        and line.start < section.end
        and line.index != section.heading_line_index
    )


def _starts_new_block(line: _Line, current: list[_Line]) -> bool:
    if not current or _BULLET_RE.match(line.text):
        return False
    if any(_BULLET_RE.match(item.text) for item in current):
        return True
    lowered = line.text.strip().casefold()
    return lowered.startswith(("project:", "project -", "certification:", "certificate:"))


def _section_blocks(section: _Section, lines: tuple[_Line, ...]) -> list[_Block]:
    blocks: list[_Block] = []
    current: list[_Line] = []

    def flush() -> None:
        if current:
            blocks.append(
                _Block(
                    lines=tuple(current),
                    start=current[0].start,
                    end=current[-1].end,
                )
            )
            current.clear()

    for line in _section_content_lines(section, lines):
        if not line.text.strip():
            flush()
        elif _starts_new_block(line, current):
            flush()
            current.append(line)
        else:
            current.append(line)
    flush()
    return blocks


def _clean_fragment(value: str, max_length: int) -> str:
    cleaned = _redact_sensitive(value)
    cleaned = _BULLET_RE.sub("", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned[:max_length].strip()


def _block_description(block: _Block, skip_lines: int = 1) -> str:
    values = []
    for line in block.lines[skip_lines:]:
        value = _clean_fragment(line.text, 500)
        if value:
            values.append(value)
    return " ".join(values)[:1200].strip()


def _technology_ids(text_value: str, normalizer: SkillNormalizer) -> list[str]:
    return sorted({match.skill_id for match in normalizer.normalize(text_value)})


def _extract_projects(
    section: _Section,
    lines: tuple[_Line, ...],
    catalog: ResumeIntelligenceCatalog,
    normalizer: SkillNormalizer,
) -> list[ProjectEvidence]:
    projects: list[ProjectEvidence] = []
    for block in _section_blocks(section, lines):
        title = _clean_fragment(block.lines[0].text, 160)
        title = re.sub(r"^(?:project|title)\s*[:\-]\s*", "", title, flags=re.IGNORECASE)
        if not title:
            continue
        block_text = " ".join(line.text for line in block.lines)
        skill_ids = _technology_ids(block_text, normalizer)
        projects.append(
            ProjectEvidence(
                project_id=f"project-{section.section_id}-{len(projects) + 1}",
                section_id=section.section_id,
                title=title[:160],
                description=_block_description(block),
                technology_skill_ids=skill_ids,
                technology_names=[
                    catalog.skills[skill_id].canonical_name for skill_id in skill_ids
                ],
                span=SourceSpan(start=block.start, end=block.end),
            )
        )
    return projects


def _split_role_organization(header: str) -> tuple[str, str]:
    at_parts = re.split(r"\s+at\s+", header, maxsplit=1, flags=re.IGNORECASE)
    if len(at_parts) == 2:
        return at_parts[0].strip(), at_parts[1].strip()
    for separator in (" | ", " — ", " – ", " - "):
        if separator in header:
            left, right = header.split(separator, 1)
            return left.strip(), right.strip()
    return header.strip(), ""


def _extract_experience(
    section: _Section,
    lines: tuple[_Line, ...],
    normalizer: SkillNormalizer,
) -> list[WorkExperienceEvidence]:
    entries: list[WorkExperienceEvidence] = []
    for block in _section_blocks(section, lines):
        if not block.lines:
            continue
        header = block.lines[0].text.strip()
        used_lines = 1
        duration_match = _DATE_RANGE_RE.search(header)
        duration = duration_match.group(0) if duration_match else ""

        if block.lines[1:]:
            second = block.lines[1].text.strip()
            second_date = _DATE_RANGE_RE.search(second)
            if second_date and len(second) <= 120:
                duration = duration or second_date.group(0)
                used_lines = 2
            elif (
                not duration
                and _COMPANY_HINT_RE.search(second)
                and len(second) <= 100
                and not _BULLET_RE.match(block.lines[1].text)
            ):
                header = f"{header} - {second}"
                used_lines = 2

        header_without_date = _DATE_RANGE_RE.sub("", header).strip(" -|,")
        role, organization = _split_role_organization(header_without_date)
        block_text = " ".join(line.text for line in block.lines)
        skill_ids = _technology_ids(block_text, normalizer)
        entries.append(
            WorkExperienceEvidence(
                experience_id=f"experience-{section.section_id}-{len(entries) + 1}",
                section_id=section.section_id,
                role=role[:160],
                organization=organization[:160],
                duration=duration[:100],
                description=_block_description(block, skip_lines=used_lines),
                skill_ids=skill_ids,
                span=SourceSpan(start=block.start, end=block.end),
            )
        )
    return entries


def _extract_education(
    section: _Section, lines: tuple[_Line, ...]
) -> list[EducationEvidence]:
    entries: list[EducationEvidence] = []
    for block in _section_blocks(section, lines):
        content = " ".join(line.text.strip() for line in block.lines).strip()
        if not content:
            continue
        degree_match = _DEGREE_RE.search(content)
        institution_match = _INSTITUTION_RE.search(content)
        year_match = _YEAR_RE.search(content)
        degree = degree_match.group(0).strip() if degree_match else ""
        if degree:
            if institution_match:
                degree = re.sub(
                    re.escape(institution_match.group(0).strip()),
                    "",
                    degree,
                    flags=re.IGNORECASE,
                )
            degree = re.split(
                r"\b(?:University|College|Institute|IIT|NIT|Academy)\b",
                degree,
                maxsplit=1,
                flags=re.IGNORECASE,
            )[0]
            degree = _YEAR_RE.sub("", degree).strip(" ,-|")
        entries.append(
            EducationEvidence(
                education_id=f"education-{section.section_id}-{len(entries) + 1}",
                section_id=section.section_id,
                institution=(
                    institution_match.group(0).strip()[:200]
                    if institution_match
                    else ""
                ),
                degree=degree[:200],
                year=(year_match.group(0).strip()[:40] if year_match else ""),
                description=_clean_fragment(content, 600),
                span=SourceSpan(start=block.start, end=block.end),
            )
        )
    return entries


def _extract_certifications(
    section: _Section, lines: tuple[_Line, ...]
) -> list[CertificationEvidence]:
    entries: list[CertificationEvidence] = []
    for block in _section_blocks(section, lines):
        name = _clean_fragment(block.lines[0].text, 200)
        name = re.sub(
            r"^(?:certification|certificate)\s*[:\-]\s*", "", name, flags=re.IGNORECASE
        )
        if not name:
            continue
        issuer = ""
        normalized_name = name
        for marker in (" by ", " from "):
            if marker in name.casefold():
                position = name.casefold().index(marker)
                normalized_name = name[:position].strip()
                issuer = name[position + len(marker):].strip()
                break
        if not issuer:
            for separator in (" | ", " - ", " — ", " – "):
                if separator in normalized_name:
                    normalized_name, issuer = normalized_name.split(separator, 1)
                    normalized_name = normalized_name.strip()
                    issuer = issuer.strip()
                    break
        entries.append(
            CertificationEvidence(
                certification_id=(
                    f"certification-{section.section_id}-{len(entries) + 1}"
                ),
                section_id=section.section_id,
                name=normalized_name[:200],
                issuer=issuer[:160],
                description=_block_description(block),
                span=SourceSpan(start=block.start, end=block.end),
            )
        )
    return entries


def _contact_metadata(
    text: str, sections: tuple[_Section, ...]
) -> ContactMetadata:
    return ContactMetadata(
        contact_section_present=any(
            section.section_type == "contact" for section in sections
        ),
        email_present=bool(_EMAIL_RE.search(text)),
        phone_present=_contains_phone(text),
        profile_link_present=bool(_PROFILE_LINK_RE.search(text)),
    )


def _analysis_warnings(
    sections: tuple[_Section, ...],
    skills: list[SkillEvidence],
    projects: list[ProjectEvidence],
    experience: list[WorkExperienceEvidence],
) -> list[str]:
    warnings: list[str] = []
    section_types = {section.section_type for section in sections}
    if not section_types.intersection(
        {
            "summary",
            "skills",
            "projects",
            "experience",
            "education",
            "certifications",
        }
    ):
        warnings.append(
            "No standard resume section headings were detected; section attribution "
            "may be limited."
        )
    if not skills:
        warnings.append("No catalog skills were detected in the resume.")
    if "projects" in section_types and not projects:
        warnings.append(
            "A projects section was detected, but no project entries could be extracted."
        )
    if "experience" in section_types and not experience:
        warnings.append(
            "An experience section was detected, but no experience entries could be extracted."
        )
    return warnings


def parse_resume_analysis(
    filename: str | None,
    content_type: str | None,
    content: bytes,
    *,
    catalog: ResumeIntelligenceCatalog | None = None,
    normalizer: SkillNormalizer | None = None,
    max_bytes: int = MAX_FILE_BYTES,
) -> ResumeAnalysis:
    """Extract a resume and return validated, evidence-backed structured data."""

    document = extract_document(
        filename, content_type, content, max_bytes=max_bytes
    )
    return analyze_extracted_document(
        document, catalog=catalog, normalizer=normalizer
    )


def analyze_extracted_document(
    document: ExtractedDocument,
    *,
    catalog: ResumeIntelligenceCatalog | None = None,
    normalizer: SkillNormalizer | None = None,
) -> ResumeAnalysis:
    """Build a validated analysis from already validated extracted text."""

    active_catalog = catalog or _catalog()
    active_normalizer = normalizer or _normalizer()
    lines = _make_lines(document.text)
    internal_sections = _parse_sections(document.text, lines)

    skills = _extract_skill_evidence(
        document.text,
        lines,
        internal_sections,
        active_catalog,
        active_normalizer,
    )
    projects: list[ProjectEvidence] = []
    experience: list[WorkExperienceEvidence] = []
    education: list[EducationEvidence] = []
    certifications: list[CertificationEvidence] = []
    for section in internal_sections:
        if section.section_type == "projects":
            projects.extend(
                _extract_projects(
                    section, lines, active_catalog, active_normalizer
                )
            )
        elif section.section_type == "experience":
            experience.extend(
                _extract_experience(section, lines, active_normalizer)
            )
        elif section.section_type == "education":
            education.extend(_extract_education(section, lines))
        elif section.section_type == "certifications":
            certifications.extend(_extract_certifications(section, lines))

    schema_sections = [
        ResumeSection(
            section_id=section.section_id,
            section_type=section.section_type,
            heading=section.heading,
            span=SourceSpan(start=section.start, end=section.end),
            line_start=section.line_start,
            line_end=section.line_end,
        )
        for section in internal_sections
    ]

    return ResumeAnalysis(
        parser_version=PARSER_VERSION,
        catalog_version=active_catalog.skills_catalog_version,
        file=FileMetadata(
            extension=document.extension,
            media_type=document.media_type,
            size_bytes=document.size_bytes,
            sha256=document.sha256,
            page_count=document.page_count,
        ),
        contact=_contact_metadata(document.text, internal_sections),
        sections=schema_sections,
        skills=skills,
        projects=projects,
        work_experience=experience,
        education=education,
        certifications=certifications,
        warnings=_analysis_warnings(
            internal_sections, skills, projects, experience
        ),
        stats=ExtractionStats(
            text_characters=len(document.text),
            section_count=len(internal_sections),
            skill_count=len(skills),
            project_count=len(projects),
            experience_count=len(experience),
            education_count=len(education),
            certification_count=len(certifications),
        ),
    )


def to_legacy_parsed_data(analysis: ResumeAnalysis) -> dict:
    """Return the existing frontend/interview shape without inventing entries."""

    return {
        "skills": [skill.canonical_name for skill in analysis.skills],
        "work_experience": [
            {
                "company": item.organization,
                "role": item.role or "Not specified",
                "duration": item.duration,
                "description": item.description,
            }
            for item in analysis.work_experience
        ],
        "projects": [
            {
                "name": item.title,
                "description": item.description,
                "technologies": item.technology_names,
            }
            for item in analysis.projects
        ],
        "education": [
            {
                "institution": item.institution,
                "degree": item.degree,
                "year": item.year,
            }
            for item in analysis.education
        ],
        "certifications": [
            {"name": item.name, "issuer": item.issuer}
            for item in analysis.certifications
        ],
    }


def build_resume_summary(analysis: ResumeAnalysis) -> str:
    """Build a deterministic, non-PII summary for existing app consumers."""

    parts = [
        f"Resume parsed with {analysis.stats.skill_count} catalog skills",
        f"{analysis.stats.project_count} project entries",
        f"{analysis.stats.experience_count} experience entries",
        f"{analysis.stats.education_count} education entries",
    ]
    summary = ", ".join(parts) + "."
    if analysis.skills:
        summary += " Skills: " + ", ".join(
            skill.canonical_name for skill in analysis.skills[:10]
        ) + "."
    if analysis.warnings:
        summary += " Extraction notes: " + " ".join(analysis.warnings)
    return summary[:1000]
